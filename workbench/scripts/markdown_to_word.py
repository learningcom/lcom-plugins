from pathlib import Path
import argparse
import hashlib
import re

from markdown_it import MarkdownIt
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


def markdown_anchor(text: str) -> str:
    """Create the same anchor format used in the Markdown glossary."""
    anchor = text.lower()
    anchor = re.sub(r"[*_`]", "", anchor)
    anchor = re.sub(r"[^\w\s-]", "", anchor)
    anchor = re.sub(r"\s+", "-", anchor)
    anchor = re.sub(r"-+", "-", anchor)
    return anchor.strip("-")


def word_bookmark_name(markdown_id: str) -> str:
    """
    Convert a Markdown anchor to a valid Word bookmark name.

    Word bookmark names cannot safely contain characters such as '-'.
    """
    name = re.sub(r"[^A-Za-z0-9_]", "_", markdown_id)
    name = "b_" + name

    # Keep bookmark names reasonably short.
    if len(name) > 40:
        digest = hashlib.sha1(
            markdown_id.encode("utf-8")
        ).hexdigest()[:7]

        name = name[:32] + "_" + digest

    return name


def plain_inline_text(token) -> str:
    """Extract visible text from an inline Markdown token."""
    if not token.children:
        return token.content

    parts = []

    for child in token.children:
        if child.type in {"text", "code_inline"}:
            parts.append(child.content)

        elif child.type in {"softbreak", "hardbreak"}:
            parts.append(" ")

    return "".join(parts)


def add_bookmark(paragraph, name: str, bookmark_id: int):
    """Add a Word bookmark around a paragraph."""
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bookmark_id))
    start.set(qn("w:name"), name)

    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bookmark_id))

    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def create_hyperlink_run(
    text: str,
    bold=False,
    italic=False,
):
    """Create the text portion of a Word hyperlink."""

    run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")

    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    run_properties.append(color)

    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_properties.append(underline)

    if bold:
        run_properties.append(OxmlElement("w:b"))

    if italic:
        run_properties.append(OxmlElement("w:i"))

    run.append(run_properties)

    text_element = OxmlElement("w:t")

    if text.startswith(" ") or text.endswith(" "):
        text_element.set(qn("xml:space"), "preserve")

    text_element.text = text
    run.append(text_element)

    return run


def add_hyperlink(
    paragraph,
    text: str,
    *,
    anchor=None,
    url=None,
    bold=False,
    italic=False,
):
    """
    Add either an internal Word hyperlink or an external URL.
    """

    hyperlink = OxmlElement("w:hyperlink")

    if anchor:
        # Internal link to a Word bookmark.
        hyperlink.set(qn("w:anchor"), anchor)
        hyperlink.set(qn("w:history"), "1")

    elif url:
        # External hyperlink.
        relationship_id = paragraph.part.relate_to(
            url,
            "http://schemas.openxmlformats.org/"
            "officeDocument/2006/relationships/hyperlink",
            is_external=True,
        )

        hyperlink.set(
            qn("r:id"),
            relationship_id,
        )

    else:
        raise ValueError(
            "Hyperlink requires either anchor or URL."
        )

    hyperlink.append(
        create_hyperlink_run(
            text,
            bold=bold,
            italic=italic,
        )
    )

    paragraph._p.append(hyperlink)


def add_inline(
    paragraph,
    inline_token,
    bookmark_map,
):
    """
    Add inline Markdown formatting to a Word paragraph.

    Supports:
    - bold
    - italics
    - internal links
    - external links
    - line breaks
    - inline code
    """

    children = inline_token.children or []

    bold = False
    italic = False
    active_href = None

    for child in children:

        if child.type == "strong_open":
            bold = True

        elif child.type == "strong_close":
            bold = False

        elif child.type == "em_open":
            italic = True

        elif child.type == "em_close":
            italic = False

        elif child.type == "link_open":
            active_href = child.attrGet("href")

        elif child.type == "link_close":
            active_href = None

        elif child.type in {"softbreak", "hardbreak"}:
            paragraph.add_run().add_break()

        elif child.type in {"text", "code_inline"}:

            text = child.content

            if active_href:

                # Internal Markdown link.
                if active_href.startswith("#"):

                    markdown_id = active_href[1:]

                    target = bookmark_map.get(
                        markdown_id
                    )

                    if target is None:
                        raise ValueError(
                            "Internal link target "
                            f"not found: {active_href}"
                        )

                    add_hyperlink(
                        paragraph,
                        text,
                        anchor=target,
                        bold=bold,
                        italic=italic,
                    )

                # External link.
                else:
                    add_hyperlink(
                        paragraph,
                        text,
                        url=active_href,
                        bold=bold,
                        italic=italic,
                    )

            else:
                run = paragraph.add_run(text)

                run.bold = bold
                run.italic = italic

                if child.type == "code_inline":
                    run.font.name = "Consolas"
                    run.font.size = Pt(9)


def add_horizontal_rule(document):
    """Add a horizontal separator."""

    paragraph = document.add_paragraph()

    paragraph_properties = (
        paragraph._p.get_or_add_pPr()
    )

    paragraph_border = OxmlElement("w:pBdr")

    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "BFBFBF")

    paragraph_border.append(bottom)
    paragraph_properties.append(
        paragraph_border
    )


def parse_table(tokens, start_index):
    """Extract a Markdown table from the token stream."""

    rows = []
    current_row = None
    current_cell = None
    header_section = False

    i = start_index + 1

    while i < len(tokens):

        token = tokens[i]

        if token.type == "table_close":
            return rows, i + 1

        elif token.type == "thead_open":
            header_section = True

        elif token.type == "thead_close":
            header_section = False

        elif token.type == "tr_open":
            current_row = []

        elif token.type == "tr_close":

            if current_row is not None:
                rows.append(current_row)

            current_row = None

        elif token.type in {
            "th_open",
            "td_open",
        }:
            current_cell = {
                "header": (
                    header_section
                    or token.type == "th_open"
                ),
                "inline": None,
            }

        elif (
            token.type == "inline"
            and current_cell is not None
        ):
            current_cell["inline"] = token

        elif token.type in {
            "th_close",
            "td_close",
        }:

            if current_row is not None:
                current_row.append(
                    current_cell
                )

            current_cell = None

        i += 1

    raise ValueError(
        "Markdown table was not closed."
    )


def add_table(
    document,
    rows,
    bookmark_map,
):
    """Create a Word table from Markdown table data."""

    if not rows:
        return

    column_count = max(
        len(row)
        for row in rows
    )

    table = document.add_table(
        rows=0,
        cols=column_count,
    )

    table.style = "Table Grid"

    for row_data in rows:

        cells = table.add_row().cells

        for column_number, cell_data in enumerate(
            row_data
        ):

            paragraph = (
                cells[column_number]
                .paragraphs[0]
            )

            inline = cell_data["inline"]

            if inline:
                add_inline(
                    paragraph,
                    inline,
                    bookmark_map,
                )

            if cell_data["header"]:
                for run in paragraph.runs:
                    run.bold = True


def build_bookmark_map(tokens):
    """
    Create the mapping:

    Markdown anchor -> Word bookmark

    Example:
    account -> b_account
    """

    bookmark_map = {}

    i = 0

    while i < len(tokens) - 2:

        if (
            tokens[i].type == "heading_open"
            and tokens[i + 1].type == "inline"
        ):

            heading = plain_inline_text(
                tokens[i + 1]
            )

            markdown_id = markdown_anchor(
                heading
            )

            bookmark_map.setdefault(
                markdown_id,
                word_bookmark_name(
                    markdown_id
                ),
            )

            i += 3

        else:
            i += 1

    return bookmark_map


def markdown_to_docx(
    input_file: str,
) -> Path:

    input_path = Path(
        input_file
    ).resolve()

    if not input_path.is_file():
        raise FileNotFoundError(
            f"Markdown file not found: "
            f"{input_path}"
        )

    output_path = (
        input_path.with_suffix(".docx")
    )

    markdown_text = (
        input_path.read_text(
            encoding="utf-8"
        )
    )

    # Parse Markdown.
    markdown = (
        MarkdownIt("commonmark")
        .enable("table")
    )

    tokens = markdown.parse(
        markdown_text
    )

    # Build Word bookmark names before
    # processing links.
    bookmark_map = (
        build_bookmark_map(tokens)
    )

    document = Document()

    document.core_properties.title = (
        input_path.stem
    )

    document.styles[
        "Normal"
    ].font.name = "Aptos"

    document.styles[
        "Normal"
    ].font.size = Pt(10.5)

    bookmark_id = 1
    list_stack = []

    i = 0

    while i < len(tokens):

        token = tokens[i]

        # -------------------------
        # Heading
        # -------------------------
        if token.type == "heading_open":

            level = int(token.tag[1])

            inline = tokens[i + 1]

            heading_text = (
                plain_inline_text(inline)
            )

            markdown_id = (
                markdown_anchor(
                    heading_text
                )
            )

            paragraph = (
                document.add_heading(
                    level=level
                )
            )

            add_inline(
                paragraph,
                inline,
                bookmark_map,
            )

            add_bookmark(
                paragraph,
                bookmark_map[markdown_id],
                bookmark_id,
            )

            bookmark_id += 1

            i += 3
            continue

        # -------------------------
        # Lists
        # -------------------------
        if token.type == "bullet_list_open":

            list_stack.append(
                "bullet"
            )

            i += 1
            continue

        if token.type == "ordered_list_open":

            list_stack.append(
                "number"
            )

            i += 1
            continue

        if token.type in {
            "bullet_list_close",
            "ordered_list_close",
        }:

            if list_stack:
                list_stack.pop()

            i += 1
            continue

        # -------------------------
        # Paragraph
        # -------------------------
        if token.type == "paragraph_open":

            inline = tokens[i + 1]

            if list_stack:

                if list_stack[-1] == "bullet":
                    style = "List Bullet"
                else:
                    style = "List Number"

                paragraph = (
                    document.add_paragraph(
                        style=style
                    )
                )

            else:
                paragraph = (
                    document.add_paragraph()
                )

            add_inline(
                paragraph,
                inline,
                bookmark_map,
            )

            i += 3
            continue

        # -------------------------
        # Table
        # -------------------------
        if token.type == "table_open":

            rows, next_index = (
                parse_table(
                    tokens,
                    i,
                )
            )

            add_table(
                document,
                rows,
                bookmark_map,
            )

            i = next_index
            continue

        # -------------------------
        # Horizontal rule
        # -------------------------
        if token.type == "hr":

            add_horizontal_rule(
                document
            )

            i += 1
            continue

        i += 1

    document.save(
        output_path
    )

    return output_path


if __name__ == "__main__":

    parser = argparse.ArgumentParser(
        description=(
            "Convert Markdown business glossary "
            "to Word DOCX."
        )
    )

    parser.add_argument(
        "input_file",
        help=(
            "Path to the Markdown file."
        ),
    )

    args = parser.parse_args()

    output = markdown_to_docx(
        args.input_file
    )

    print(
        f"Created: {output}"
    )